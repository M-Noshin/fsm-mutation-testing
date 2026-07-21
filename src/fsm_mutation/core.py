"""Shared search primitives: kill semantics, metrics, and result reporting.

Every algorithm in this package uses the functions defined here. That is the
whole point of the module: if BFS, DFS, Q-learning and look-ahead all step the
mutant population through *the same* code, then differences in their reported
scores are differences between the algorithms and nothing else.

The kill condition
------------------
A mutant ``M`` is killed by an input sequence ``s`` if and only if applying
``s`` from the initial state to the specification ``S`` and to ``M`` yields
**different output sequences**.

This is the definition used in the paper ("the output sequences produced by S
and M to the input sequence are different") and the one implemented by the
generator's ``distinguish_machines``. It is *not* the same as "the mutant moved
to a different state": a transfer fault produces no observable difference at
the step where it is traversed, and is only detected later -- if ever -- when
the divergent state emits a different output. ``KillCondition.OUTPUT_OR_STATE``
reproduces the older, weaker behaviour for backward comparison, but it
overstates what a test can actually detect and should not be used for new
results.
"""

from __future__ import annotations

import enum
import sys
import time
from dataclasses import dataclass, field

import numpy as np
import psutil

from .io import UNDEFINED, MutantSet, Spec

#: Checkpoints (seconds) at which metrics are sampled.
INTERVALS = (30, 60, 90, 120, 150, 180, 240, 300, 360)
INTERVAL_NAMES = ("30s", "1min", "1min30s", "2min", "2min30s",
                  "3min", "4min", "5min", "6min")


class KillCondition(enum.Enum):
    #: Correct: only an output difference is observable.
    OUTPUT = "output"
    #: Legacy: also treats a next-state difference as a kill. Overstates scores.
    OUTPUT_OR_STATE = "output_or_state"


# --------------------------------------------------------------------------
# Population state
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class Survivors:
    """The still-alive mutants and the state each of them currently occupies.

    Only survivors are stored, so the footprint shrinks as the search kills
    mutants -- unlike a fixed-length array padded with ``None``, whose size is
    constant and dominates every memory measurement.
    """

    idx: np.ndarray    # int32, mutant ids
    state: np.ndarray  # int32, current state of each

    @property
    def count(self) -> int:
        return int(self.idx.size)

    def nbytes(self) -> int:
        return int(self.idx.nbytes + self.state.nbytes)

    def key(self) -> bytes:
        """Hashable identity, for de-duplicating search nodes."""
        return self.idx.tobytes() + b"|" + self.state.tobytes()


def initial_survivors(mutants: MutantSet, spec: Spec) -> Survivors:
    return Survivors(
        idx=np.arange(mutants.n_mutants, dtype=np.int32),
        state=np.full(mutants.n_mutants, spec.initial, dtype=np.int32),
    )


def step(spec: Spec, mutants: MutantSet, spec_state: int, surv: Survivors,
         inp: int, kill: KillCondition = KillCondition.OUTPUT):
    """Apply one input to the specification and to every surviving mutant.

    Returns ``(next_spec_state, new_survivors, killed_ids)`` where
    ``killed_ids`` is the array of mutant ids killed *by this step*.

    Raises ``ValueError`` if the input is undefined at ``spec_state``; callers
    must only choose inputs from ``spec.defined_inputs``.
    """
    sp_next = int(spec.next_state[spec_state, inp])
    sp_out = int(spec.output[spec_state, inp])
    if sp_next == UNDEFINED:
        raise ValueError(
            f"input {inp} is undefined at specification state {spec_state}"
        )

    if surv.count == 0:
        return sp_next, surv, np.empty(0, dtype=np.int32)

    # Default behaviour of every mutant is the specification's behaviour;
    # the sparse delta overrides it only where the mutant actually differs.
    m_next = spec.next_state[surv.state, inp].astype(np.int32, copy=True)
    m_out = spec.output[surv.state, inp].astype(np.int32, copy=True)

    if mutants.keys.size:
        keys = (surv.idx.astype(np.int64) * mutants.stride
                + surv.state.astype(np.int64) * spec.n_inputs + inp)
        pos = np.searchsorted(mutants.keys, keys)
        np.clip(pos, 0, mutants.keys.size - 1, out=pos)
        hit = mutants.keys[pos] == keys
        if hit.any():
            m_next[hit] = mutants.delta_next[pos[hit]]
            m_out[hit] = mutants.delta_out[pos[hit]]

    # A mutant with no transition on this input produces no output where the
    # specification produces one: an observable difference, hence killed.
    killed = (m_next == UNDEFINED) | (m_out != sp_out)
    if kill is KillCondition.OUTPUT_OR_STATE:
        killed |= (m_next != sp_next)

    alive = ~killed
    return (
        sp_next,
        Survivors(idx=surv.idx[alive], state=m_next[alive]),
        surv.idx[killed],
    )


def kills_of_sequence(spec: Spec, mutants: MutantSet, seq,
                      kill: KillCondition = KillCondition.OUTPUT) -> int:
    """Number of mutants killed by ``seq`` applied from the initial state."""
    s = spec.initial
    surv = initial_survivors(mutants, spec)
    for inp in seq:
        s, surv, _ = step(spec, mutants, s, surv, int(inp), kill)
    return mutants.n_mutants - surv.count


# --------------------------------------------------------------------------
# Results
# --------------------------------------------------------------------------

@dataclass
class SearchResult:
    """What a search produced.

    Two distinct artefacts are reported, because tree searches naturally yield
    a *suite* of sequences while Q-learning and look-ahead yield a *single*
    sequence. Scoring one on the other's metric is what made earlier results
    incomparable, so both are always recorded for every algorithm.
    """

    algorithm: str

    # Best single distinguishing test found.
    best_sequence: list = field(default_factory=list)
    best_kills: int = 0

    # Greedy suite: every sequence that contributed at least one new kill.
    suite_size: int = 0
    suite_cost: int = 0          # total number of inputs across the suite
    suite_kills: int = 0

    nodes: int = 0
    max_depth: int = 0
    episodes: int = 0
    elapsed: float = 0.0
    terminated_at: float | None = None
    stop_reason: str = ""
    peak_memory_mb: float = float("nan")

    @property
    def ds_length(self) -> int:
        """Length of the best single distinguishing test."""
        return len(self.best_sequence)


class Tracker:
    """Accumulates the two artefacts and samples metrics on a wall clock.

    ``union`` is a boolean mask over all mutants recording everything killed
    anywhere in the search; it backs the suite metrics. ``best_*`` records the
    single best sequence. Both are updated incrementally, so tracking costs
    O(mutants killed at this step) rather than O(total mutants).
    """

    def __init__(self, spec: Spec, mutants: MutantSet, algorithm: str,
                 budget: float, kill: KillCondition = KillCondition.OUTPUT):
        self.spec = spec
        self.mutants = mutants
        self.total = mutants.n_mutants
        self.kill = kill
        self.budget = budget
        self.result = SearchResult(algorithm=algorithm)

        self._union = np.zeros(self.total, dtype=bool)
        self._union_count = 0
        self._proc = psutil.Process()
        self._start = time.perf_counter()
        self._next_interval = 0

        n = len(INTERVALS)
        self.metrics = {
            "mutation_score": [np.nan] * n,
            "suite_score": [np.nan] * n,
            "memory_mb": [np.nan] * n,
            "nodes": [np.nan] * n,
            "ds_length": [np.nan] * n,
            "max_depth": [np.nan] * n,
            "survivors": [np.nan] * n,
        }

    # -- clock ------------------------------------------------------------

    @property
    def elapsed(self) -> float:
        return time.perf_counter() - self._start

    def expired(self) -> bool:
        return self.elapsed > self.budget

    def checkpoint(self) -> None:
        """Sample metrics for every interval boundary that has just passed."""
        t = self.elapsed
        while (self._next_interval < len(INTERVALS)
               and t >= INTERVALS[self._next_interval]):
            self._record(self._next_interval)
            self._next_interval += 1

    def _record(self, i: int) -> None:
        r = self.result
        self.metrics["mutation_score"][i] = self._pct(r.best_kills)
        self.metrics["suite_score"][i] = self._pct(self._union_count)
        self.metrics["memory_mb"][i] = self._proc.memory_info().rss / (1024 ** 2)
        self.metrics["nodes"][i] = r.nodes
        self.metrics["ds_length"][i] = r.ds_length
        self.metrics["max_depth"][i] = r.max_depth
        self.metrics["survivors"][i] = self.total - self._union_count

    def _pct(self, killed: int) -> float:
        return (killed / self.total * 100.0) if self.total else 0.0

    def finish(self, reason: str) -> SearchResult:
        """Close out the run.

        Intervals after termination are left as NaN rather than back-filled.
        A run that ended at 100 s has no 3-minute measurement, and writing the
        RSS observed at 100 s into that cell produces a memory curve that
        appears to collapse. ``Terminated At`` carries the real information.
        """
        import resource

        r = self.result
        r.elapsed = self.elapsed
        r.terminated_at = r.elapsed
        r.stop_reason = reason
        r.suite_kills = self._union_count
        # True high-water mark for the process. Only meaningful when the run
        # has a process to itself -- see the note on --isolate in run.py.
        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        # Linux reports kilobytes, macOS reports bytes.
        r.peak_memory_mb = peak / 1024 if sys.platform != "darwin" else peak / 1024 ** 2
        # Intervals already elapsed but not yet sampled still get a reading.
        while (self._next_interval < len(INTERVALS)
               and INTERVALS[self._next_interval] <= r.elapsed):
            self._record(self._next_interval)
            self._next_interval += 1
        return r

    # -- accumulation -----------------------------------------------------

    def observe_step(self, killed_ids: np.ndarray, path_len: int) -> int:
        """Record the kills made by one edge of the search.

        Returns the number of *newly* killed mutants (not killed anywhere
        earlier in this run). If the edge contributed anything new, the path
        that reaches it is added to the greedy suite.
        """
        if killed_ids.size == 0:
            return 0
        fresh = killed_ids[~self._union[killed_ids]]
        if fresh.size == 0:
            return 0
        self._union[fresh] = True
        self._union_count += int(fresh.size)
        self.result.suite_size += 1
        self.result.suite_cost += path_len
        return int(fresh.size)

    def observe_path(self, sequence, kills: int) -> bool:
        """Offer a complete path as a candidate best single test.

        A path wins if it kills more mutants, or kills the same number with a
        shorter sequence. Returns True if it was adopted.
        """
        r = self.result
        better = kills > r.best_kills
        shorter = (kills == r.best_kills and kills > 0
                   and len(sequence) < len(r.best_sequence))
        if better or shorter:
            r.best_kills = kills
            r.best_sequence = list(sequence)
            return True
        return False

    def note_node(self, depth: int, n: int = 1) -> None:
        self.result.nodes += n
        if depth > self.result.max_depth:
            self.result.max_depth = depth


# --------------------------------------------------------------------------
# Row assembly
# --------------------------------------------------------------------------

def result_row(spec: Spec, mutants: MutantSet, result: SearchResult,
               tracker: Tracker, *, machine_class: str, mutant_file: str,
               max_depth_cap, seed: int, kill: KillCondition) -> dict:
    """Flatten a run into the one row that gets appended to the raw workbook."""
    row = {
        "Machine #": spec.fsm_id,
        "Machine Class": machine_class,
        "# states": spec.n_states,
        "# inputs": spec.n_inputs,
        "# outputs": spec.n_outputs,
        "# transitions": spec.n_transitions,
        "Algorithm": result.algorithm,
        "Number of Mutants": mutants.n_mutants,
        "Mutant File": mutant_file,
        "Max search depth": max_depth_cap if max_depth_cap else "none",
        "Kill Condition": kill.value,
        "Seed": seed,

        "Final Mutation Score (%)": tracker._pct(result.best_kills),
        "Final Killed Mutants": result.best_kills,
        "Final DS Length": result.ds_length,
        "Best Sequence": spec.render(result.best_sequence),

        "Suite Score (%)": tracker._pct(result.suite_kills),
        "Suite Killed Mutants": result.suite_kills,
        "Suite Size (# tests)": result.suite_size,
        "Suite Cost (# inputs)": result.suite_cost,

        "Total Nodes": result.nodes,
        "Max Depth Reached": result.max_depth,
        "Episodes": result.episodes,
        "Elapsed Time (sec)": result.elapsed,
        "Terminated At (sec)": result.terminated_at,
        "Stop Reason": result.stop_reason,
        "Peak Memory (MB)": result.peak_memory_mb,
        "Mutant Storage (MB)": mutants.nbytes() / (1024 ** 2),
    }
    labels = {
        "mutation_score": "Mutation Score after {} (%)",
        "suite_score": "Suite Score after {} (%)",
        "memory_mb": "Memory Usage (MB) after {}",
        "nodes": "Total # nodes after {}",
        "ds_length": "DS Length after {}",
        "max_depth": "Max Depth Reached after {}",
        "survivors": "Surviving Mutants after {}",
    }
    for metric, template in labels.items():
        for name, value in zip(INTERVAL_NAMES, tracker.metrics[metric]):
            row[template.format(name)] = value
    return row


def append_row(path, row: dict, retries: int = 5, delay: float = 1.0) -> None:
    """Append one result row to an xlsx workbook, creating it if needed.

    Writes via a temporary file and an atomic replace, and retries when the
    workbook is open in Excel.
    """
    import os
    import tempfile
    from openpyxl import Workbook, load_workbook

    path = str(path)
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)

    for attempt in range(retries + 1):
        try:
            if not os.path.exists(path):
                wb = Workbook()
                ws = wb.active
                ws.title = "results"
                ws.append(list(row.keys()))
                ws.append(list(row.values()))
            else:
                wb = load_workbook(path)
                ws = wb.active
                header = [c.value for c in ws[1]]
                if header != list(row.keys()):
                    # Schema drift would silently misalign columns.
                    missing = set(row) - set(header)
                    extra = set(header) - set(row)
                    raise ValueError(
                        f"column mismatch in {path}: "
                        f"missing={sorted(missing)} unexpected={sorted(extra)}"
                    )
                ws.append([row[k] for k in header])

            fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=parent or None)
            os.close(fd)
            wb.save(tmp)
            wb.close()
            os.replace(tmp, path)
            return
        except PermissionError:
            if attempt == retries:
                raise
            time.sleep(delay)
